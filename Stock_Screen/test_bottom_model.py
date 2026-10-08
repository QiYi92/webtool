import copy
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from zoneinfo import ZoneInfo

import numpy as np
import pandas as pd
import requests
from openpyxl import load_workbook

import screen_bowl_shape as screen
from bottom_model import completed_daily_bars, evaluate_bottom
from excel_exporter import export_sector_summary_excel


def sample(strong=False):
    close = np.r_[np.linspace(10, 6.3, 205), 6.3 + .08*np.sin(np.arange(40)*np.pi/5), np.full(5, 6.3)]
    if strong:
        close[-5:] = [6.34, 6.38, 6.42, 6.47, 6.52]
    df = pd.DataFrame({"日期": pd.bdate_range(end="2026-09-07", periods=250), "开盘":close,
                       "收盘":close, "最高":close+.05, "最低":close-.05, "成交量":np.full(250,100.)})
    if strong:
        df.loc[245:, "成交量"] = 130.
    return df


class BottomModelTests(unittest.TestCase):
    def setUp(self):
        self.config = screen.load_screen_config(str(Path(__file__).parent / "configs/default_bottom.json"))

    def test_both_stages(self):
        for strong, stage in [(False,"横盘未突破"),(True,"出现转强")]:
            passed, metrics = evaluate_bottom(sample(strong), self.config)
            self.assertTrue(passed, metrics)
            self.assertEqual(metrics["碗型阶段"], stage)
            self.assertEqual(metrics["筛选日期"], "2026-09-07")

    def test_strength_boundaries_only_change_stage(self):
        data = sample(True)
        _, metrics = evaluate_bottom(data, self.config)
        for key, boundary in (("strength_volume_ratio", metrics["近期量比"]), ("strength_min_gain", metrics["近期涨幅"])):
            config = copy.deepcopy(self.config)
            config["bottom"][key] = boundary
            self.assertEqual(evaluate_bottom(data, config)[1]["碗型阶段"], "出现转强")
            config["bottom"][key] += 1e-8
            passed, result = evaluate_bottom(data, config)
            self.assertTrue(passed)
            self.assertEqual(result["碗型阶段"], "横盘未突破")

    def test_rejections(self):
        cases = {}
        high = sample(); high.loc[:204, ["开盘","收盘","最高","最低"]] *= .64
        cases["高位"] = high
        falling = sample(); prices=np.linspace(6.9,5.9,40)
        for col, offset in [("收盘",0),("开盘",0),("最高",.05),("最低",-.05)]:
            falling.loc[205:244,col] = prices+offset
        cases["阴跌"] = falling
        broken = sample(); broken.loc[245, ["开盘","收盘","最高","最低"]] = [6,6,6.05,5.95]
        cases["破位"] = broken
        rebound = sample(); rebound.loc[249,["开盘","收盘","最高","最低"]]=[7.4,7.4,7.45,7.35]
        cases["反弹过大"] = rebound
        cases["数据不足"] = sample().iloc[1:]
        missing=sample(); missing.loc[249,"收盘"]=np.nan; cases["缺失"] = missing
        zero=sample(); zero.loc[220,"成交量"]=0; cases["零量"] = zero
        constant=sample(); constant[["开盘","收盘","最高","最低"]]=6.3; cases["零区间"] = constant
        for name, data in cases.items():
            with self.subTest(name=name):
                passed, metrics=evaluate_bottom(data,self.config)
                self.assertFalse(passed)
                self.assertTrue(metrics["未通过原因"])

    def test_bottom_model_uses_akshare_qfq_and_cutoff(self):
        with patch.object(screen, "fetch_stock_hist_em", return_value=sample()) as history:
            result = screen.fetch_bottom_k("002531", 260, "2026-09-04")
        self.assertEqual(result["日期"].max(), pd.Timestamp("2026-09-04"))
        self.assertEqual(history.call_args.kwargs["adjust"], "qfq")

    def test_bottom_model_does_not_fallback_when_akshare_fails(self):
        with patch.object(screen, "fetch_stock_hist_em", side_effect=screen.MarketDataConnectionError("AKShare offline")):
            with self.assertRaises(screen.MarketDataConnectionError):
                screen.fetch_bottom_k("605599", 260, "2026-09-07")

    def test_daily_history_uses_akshare_adapter(self):
        payload = pd.DataFrame([{"日期":"2026-09-23", "开盘":100, "收盘":101, "最高":102,
                                 "最低":99, "成交量":1000, "成交额":100000, "换手率":0.5,
                                 "股票代码":"600519"}])
        with patch.object(screen, "fetch_stock_history", return_value=payload) as fetch:
            bars = screen.fetch_stock_hist_em("600519", "20260923", "20260923", adjust="qfq")
        self.assertEqual(fetch.call_args.args, ("600519", "20260923", "20260923", "qfq", "daily"))
        self.assertEqual(len(bars), 1)
        self.assertEqual(bars.iloc[0]["收盘"], 101)

    def test_long_history_uses_akshare_and_preserves_liquidity_fields(self):
        dates = pd.bdate_range("2021-09-24", "2026-09-23")
        raw = pd.DataFrame({"日期": dates, "开盘":10., "收盘":10., "最高":10.1, "最低":9.9,
                            "成交量":100., "换手率":0.15, "成交额":1234000.})
        with patch.object(screen, "fetch_stock_history", return_value=raw) as fetch:
            result = screen.fetch_akshare_qfq_history("600000", "2021-09-24", "2026-09-23")
        self.assertEqual(fetch.call_args.args, ("600000", "2021-09-24", "2026-09-23", "qfq", "daily"))
        self.assertEqual(len(result), len(dates))
        self.assertAlmostEqual(result.iloc[-1]["换手率"], .15)
        self.assertAlmostEqual(result.iloc[-1]["成交额"], 1234000)

    def test_threshold_boundaries(self):
        df=sample()
        _, metrics=evaluate_bottom(df,self.config)
        for parameter, metric, minimum in [("max_year_position","年内位置",False),("min_drawdown","年内回撤",True),
                                           ("max_platform_amplitude","平台振幅",False),("max_close_amplitude","平台收盘振幅",False)]:
            with self.subTest(parameter=parameter):
                config=copy.deepcopy(self.config)
                config["bottom"][parameter]=metrics[metric]
                self.assertTrue(evaluate_bottom(df,config)[0])
                config["bottom"][parameter]+=1e-8 if minimum else -1e-8
                self.assertFalse(evaluate_bottom(df,config)[0])

    def test_completed_bars_and_no_future_data(self):
        df=sample()
        now=datetime(2026,9,7,14,59,tzinfo=ZoneInfo("Asia/Shanghai"))
        self.assertEqual(len(completed_daily_bars(df,now)),249)
        self.assertEqual(len(completed_daily_bars(df,now.replace(hour=15))),250)
        future=df.iloc[[-1]].copy(); future["日期"]=pd.Timestamp("2026-09-08")
        combined=pd.concat([df,future],ignore_index=True)
        self.assertEqual(len(completed_daily_bars(combined,now.replace(hour=15))),250)

    def test_logs_match_bowl_style_and_report_funnel_counts(self):
        universe = pd.DataFrame([
            {"代码": f"00000{i}", "名称": f"测试{i}", "总市值": 20e9,
             "上市时间": None if i == 0 else 20100101, "板块": "测试板块", "概念": ""}
            for i in range(6)
        ])
        broken = sample()
        broken.loc[245, ["开盘", "收盘", "最高", "最低"]] = [6, 6, 6.05, 5.95]
        responses = [sample(), sample(True), sample().iloc[-45:], broken, RuntimeError("行情失败")]
        with (
            patch.object(screen, "get_stock_universe", return_value=universe),
            patch.object(screen, "fetch_bottom_k", side_effect=responses),
            patch.object(screen, "check_sideways_candidate", return_value=(True, {"规则版本": "sideways_v2"})),
            self.assertLogs(level="INFO") as captured,
        ):
            result = screen.screen_stocks(config=self.config, sleep_seconds=0)
        self.assertEqual(len(result), 2)
        messages = [record.getMessage() for record in captured.records]
        self.assertIn("命中: 000001 测试1", messages)
        self.assertIn("命中: 000002 测试2", messages)
        self.assertIn("进度: 1/6", messages)
        self.assertIn("========== 筛选统计 ==========", messages)
        self.assertEqual(messages[-1], "========== 统计结束 ==========")
        for key, count in {
            "基础过滤后待扫描": 6, "上市时间不足/缺失": 1, "上市时间通过": 5,
            "K线数据不足": 1, "K线数据通过": 3, "底部平台不通过": 1,
            "底部平台通过": 2, "横盘未突破": 1, "出现转强": 1,
            "接口异常/其他异常": 1, "最终命中": 2,
        }.items():
            self.assertIn(f"{key}: {count}", messages)
        self.assertFalse(any(message.startswith(("跳过 000003", "跳过 000004")) for message in messages))
        self.assertIn("跳过 000005 测试5: 行情失败", messages)
        self.assertFalse(any("一年趋势" in message or "碗型" in message for message in messages))

    def test_market_data_connection_failure_stops_scan_immediately(self):
        universe = pd.DataFrame([{
            "代码": "000001", "名称": "测试", "总市值": 20e9,
            "上市时间": 20100101, "板块": "测试板块", "概念": "",
        }])
        try:
            raise RuntimeError("行情接口重试失败") from requests.ConnectionError("offline")
        except RuntimeError as connection_failure:
            with patch.object(screen, "get_stock_universe", return_value=universe), \
                    patch.object(screen, "fetch_bottom_k", side_effect=connection_failure):
                with self.assertRaisesRegex(screen.MarketDataConnectionError, "任务已停止"):
                    screen.screen_stocks(config=self.config, sleep_seconds=0)

    def test_screen_export_and_mandatory_checks(self):
        universe=pd.DataFrame([{"代码":"002531","名称":"测试","总市值":20e9,"上市时间":20100101,"板块":"测试板块","概念":""}])
        config=copy.deepcopy(self.config)
        config["basic"].update(enable_trend_filter=True,enable_volume_filter=True,enable_bowl_filter=False,enable_k_data_filter=False)
        with patch.object(screen,"get_stock_universe",return_value=universe), patch.object(screen,"fetch_bottom_k",return_value=sample(True)), patch.object(screen,"is_one_year_uptrend",side_effect=AssertionError("旧过滤不应调用")):
            result=screen.screen_stocks(config=config,sleep_seconds=0)
        self.assertEqual(result.iloc[0]["碗型阶段"],"出现转强")
        with tempfile.TemporaryDirectory() as directory:
            path=Path(directory)/"result.xlsx"
            export_sector_summary_excel(result,path)
            wb=load_workbook(path)
            self.assertIn("横盘未突破",wb.sheetnames)
            self.assertIn("出现转强",wb.sheetnames)
            self.assertIn("年内位置",[cell.value for cell in wb["模型指标"][1]])
            wb.close()
        with patch.object(screen,"get_stock_universe",return_value=universe), patch.object(screen,"fetch_bottom_k",return_value=sample().iloc[-45:]):
            self.assertTrue(screen.screen_stocks(config=config,sleep_seconds=0).empty)

if __name__ == "__main__":
    unittest.main()
